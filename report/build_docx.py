# -*- coding: utf-8 -*-
"""Build the BTL report docx (Vietnamese, Times New Roman 13, A4) with python-docx."""
import os
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "img")
import sys
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "LTM_Nhom_GameDuaXe_Java.docx")

FONT = "Times New Roman"
BODY_PT = 13
CODE_FONT = "Consolas"
HEADER_FILL = "DCE6F1"
CODE_FILL = "F3F3F3"
CONTENT_WIDTH_IN = 6.27  # A4 minus 1in margins

doc = Document()

# ---------------------------------------------------------------- page setup
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Inches(1))


def set_font(run, name=FONT, size=BODY_PT, bold=None, italic=None, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def style_base(style, size=BODY_PT, bold=False, color="000000"):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    # drop theme fonts so Word does not override
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]


normal = doc.styles["Normal"]
style_base(normal)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

for name, size in (("Heading 1", 14), ("Heading 2", 13), ("Heading 3", 13)):
    st = doc.styles[name]
    style_base(st, size=size, bold=True)
    st.font.italic = name == "Heading 3"
    st.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 8)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True

for name in ("List Bullet", "List Number"):
    style_base(doc.styles[name])
    doc.styles[name].paragraph_format.space_after = Pt(3)

# ---------------------------------------------------------------- helpers
fig_no = [0]
tab_no = [0]


def heading(text, level=1):
    return doc.add_heading(text, level=level)


def para(text, align="justify", bold=False, italic=False, size=BODY_PT, space_after=6, indent=True):
    p = doc.add_paragraph()
    if align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "left":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if indent and align == "justify":
        p.paragraph_format.first_line_indent = Cm(1.0)
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, italic=italic)
    return p


def rich(parts, align="justify", indent=True):
    """parts: list of (text, bold) tuples."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if align == "justify" else WD_ALIGN_PARAGRAPH.LEFT
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.0)
    for text, b in parts:
        set_font(p.add_run(text), bold=b)
    return p


def bullets(items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.left_indent = Cm(1.25)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if isinstance(it, tuple):
            set_font(p.add_run(it[0]), bold=True)
            set_font(p.add_run(it[1]))
        else:
            set_font(p.add_run(it))


def shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def cell_margins(table, top=60, bottom=60, left=100, right=100):
    tblpr = table._element.tblPr
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblpr.append(mar)


def fix_layout(t, widths_in):
    tblpr = t._element.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:type"), "dxa")
    tblw.set(qn("w:w"), str(int(sum(widths_in) * 1440)))
    for i, w in enumerate(widths_in):
        t.columns[i].width = Inches(w)


def no_split(row):
    trpr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:cantSplit")
    e.set(qn("w:val"), "true")
    trpr.append(e)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trpr.append(e)


def table(headers, rows, widths_in, caption=None, size=12, first_col_bold=False, align_center_cols=()):
    if caption:
        tab_no[0] += 1
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        set_font(p.add_run(f"Bảng {tab_no[0]}. {caption}"), bold=True, size=12)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell_margins(t)
    fix_layout(t, widths_in)
    hdr = t.rows[0]
    repeat_header(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        c.width = Inches(widths_in[i])
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(c, HEADER_FILL)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(h), bold=True, size=size)
    for r in rows:
        row = t.add_row()
        no_split(row)
        cells = row.cells
        for i, val in enumerate(r):
            c = cells[i]
            c.width = Inches(widths_in[i])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if i in align_center_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            lines = val if isinstance(val, list) else [val]
            for k, ln in enumerate(lines):
                if k > 0:
                    p = c.add_paragraph()
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                set_font(p.add_run(ln), size=size, bold=(first_col_bold and i == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(filename, caption, width_in=CONTENT_WIDTH_IN):
    fig_no[0] += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(os.path.join(IMG, filename), width=Inches(width_in))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(8)
    set_font(c.add_run(f"Hình {fig_no[0]}. {caption}"), italic=True, size=12)


def code(lines, size=9):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell_margins(t, top=80, bottom=80, left=140, right=140)
    fix_layout(t, [CONTENT_WIDTH_IN])
    no_split(t.rows[0])
    c = t.rows[0].cells[0]
    c.width = Inches(CONTENT_WIDTH_IN)
    shade(c, CODE_FILL)
    p = c.paragraphs[0]
    for i, ln in enumerate(lines):
        if i > 0:
            p = c.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        set_font(p.add_run(ln if ln else " "), name=CODE_FONT, size=size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    set_font(r, size=11)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), tag)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = text
        r._r.append(e)


add_page_number(sec)

# ================================================================ TITLE
for txt, sz in (("BÀI TẬP LỚN", 16), ("MÔN LẬP TRÌNH MẠNG", 14)):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(txt), bold=True, size=sz)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(4)
set_font(p.add_run("GAME ĐUA XE THI ĐẤU ĐỐI KHÁNG ONLINE"), bold=True, size=18)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(16)
set_font(p.add_run("Bài tập 2 – Mô tả kiến trúc, luồng và thiết kế hệ thống"), bold=True, size=13)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
set_font(p.add_run("Lớp: LTM-2026-2-N01 · Ngôn ngữ Java · TCP Socket · Java Swing · MySQL (JDBC)"), italic=True, size=13)
table(["STT", "Thành viên", "Phần phụ trách"], [
    ["1", "Phạm Thị Thu Phương", "Server: kết nối, phiên, lời mời, phòng đua, trọng tài (mục 4; task S1–S7)"],
    ["2", "Vũ Văn Hiếu", "Client sảnh: mạng phía client, đăng nhập, sảnh chờ, lời mời (mục 5; task C1–C7)"],
    ["3", "Trần Việt Anh", "Nền tảng (Maven, common, schema, DAO, kiểm thử, đóng gói) và màn hình đua (mục 6; task D1–D7, R1–R6)"],
    ["4", "Nguyễn Trần Mai Anh", "Dữ liệu SQL: bảng xếp hạng, đăng ký, lịch sử trận, ghi diễn biến (mục 7; task M1–M6)"],
], [0.6, 2.2, 3.47], caption="Danh sách thành viên nhóm", align_center_cols=(0,))

# ================================================================ 1. BÀI TOÁN
heading("1. Mô tả bài toán và yêu cầu")
para("Đề tài xây dựng một trò chơi đua xe thi đấu đối kháng trực tuyến theo mô hình Client – Server. "
     "Hệ thống có một server và nhiều client. Server lưu toàn bộ thông tin người chơi, dữ liệu trận đấu và bảng "
     "xếp hạng; client là chương trình Java Swing để người chơi đăng nhập, xem danh sách người chơi đang online, "
     "thách đấu và điều khiển xe trong trận. Toàn bộ luật thắng thua, thời gian chờ và việc đồng bộ trạng thái "
     "hai xe do server quyết định, client chỉ hiển thị và gửi thao tác của người chơi.")
para("Các yêu cầu chức năng được tách từ đề bài và đánh mã để tham chiếu trong các phần thiết kế phía sau:", indent=True)

table(["Mã", "Yêu cầu", "Mô tả chi tiết"], [
    ["YC01", "Đăng nhập", "Người chơi đăng nhập từ máy client bằng tài khoản đã đăng ký; server kiểm tra với CSDL."],
    ["YC02", "Danh sách online", "Sau khi đăng nhập, client hiển thị danh sách người chơi đang online gồm tên, tổng điểm và trạng thái (Rảnh / Đang thi đấu)."],
    ["YC03", "Thách đấu", "Người chơi click vào tên đối thủ đang Rảnh trong danh sách để gửi lời mời thách đấu."],
    ["YC04", "Phản hồi lời mời", "Người bị thách đấu chọn Chấp nhận (OK) hoặc Từ chối (Reject); sau 30 giây không phản hồi lời mời tự hủy."],
    ["YC05", "Tạo phòng đua", "Khi chấp nhận, hai người được đưa vào một phòng đua; server làm trọng tài; đếm ngược 3, 2, 1 trước khi bắt đầu."],
    ["YC06", "Đường đua", "Mỗi người điều khiển một xe trên đường đua riêng có cùng chiều dài và chướng ngại vật giống nhau; chướng ngại vật là các xe cộ chạy cùng chiều, người chơi phải đổi làn để vượt. Hai đường đua hiển thị song song trên màn hình."],
    ["YC07", "Điều khiển và đồng bộ", "Người chơi điều khiển xe bằng bàn phím; client liên tục gửi trạng thái xe lên server; server đồng bộ và cập nhật vị trí hai xe cho cả hai người chơi."],
    ["YC08", "Xác định kết quả", "Người về đích trước thắng: thắng 1 điểm, thua 0 điểm. Về đích cùng lúc tính hòa, mỗi người 1 điểm."],
    ["YC09", "Thi đấu tiếp", "Sau mỗi trận server hỏi hai người có muốn thi đấu tiếp; cả hai đồng ý thì tạo ván mới, một người từ chối thì kết thúc."],
    ["YC10", "Thoát trận", "Người chơi có thể thoát khỏi trận bất cứ lúc nào; hệ thống báo cho người còn lại, xử thua người thoát và kết thúc trận."],
    ["YC11", "Mất kết nối", "Server phát hiện client mất mạng hoặc tắt ứng dụng sau khoảng timeout, xử thua người đó và thông báo cho đối thủ."],
    ["YC12", "Bảng xếp hạng", "Kết quả các trận được lưu tại server; mọi người chơi xem được bảng xếp hạng sắp xếp theo tổng điểm giảm dần rồi tổng số trận thắng giảm dần."],
], [0.7, 1.45, 4.12], caption="Yêu cầu chức năng của hệ thống", first_col_bold=True, align_center_cols=(0,))

para("Ngoài các yêu cầu chức năng, hệ thống phải đáp ứng một số yêu cầu phi chức năng: viết bằng Java, dùng "
     "TCP Socket để bảo đảm thông điệp đến đúng thứ tự và không mất; server phục vụ đồng thời nhiều client bằng "
     "đa luồng; dữ liệu lâu dài lưu trong MySQL; độ trễ đồng bộ vị trí xe đủ nhỏ để hai người chơi thấy trận đấu "
     "diễn ra mượt mà (khoảng 50 ms một lần cập nhật).")

# ================================================================ 2. KIẾN TRÚC CHUNG
heading("2. Kiến trúc chung của hệ thống")
para("Hệ thống sử dụng kiến trúc Client – Server. Nhiều client Java Swing kết nối đến một Game Server qua TCP "
     "Socket. Server là nơi duy nhất quản lý người chơi online, lời mời thách đấu, phòng đua, luật thắng thua, "
     "trạng thái kết nối và kết quả trận. MySQL lưu dữ liệu lâu dài về người chơi, trận đấu và các sự kiện trong trận.")
para("Mỗi client được tổ chức theo mô hình MVC: View hiển thị giao diện Swing, Controller nhận thao tác của người "
     "chơi và giao tiếp mạng, Model chứa dữ liệu Player, Match, RaceState và RaceEvent. Server dùng ServerSocket "
     "để nhận kết nối; mỗi client được xử lý bởi một ClientHandler chạy trên luồng riêng để nhiều người chơi có thể "
     "hoạt động đồng thời. Thông điệp trao đổi là đối tượng Message (kiểu, dữ liệu kèm theo) được gửi qua "
     "ObjectOutputStream và nhận bằng ObjectInputStream.")
figure("fig1_architecture.png", "Kiến trúc tổng thể Client – Server theo MVC")

table(["Thành phần trên hình", "Chức năng", "Mô tả"], [
    ["Client A, B, … N", "Giao diện và gửi sự kiện", "Đăng nhập, xem sảnh, thách đấu, điều khiển xe và hiển thị dữ liệu trận. Mỗi client là một tiến trình Java Swing độc lập."],
    ["Mũi tên Client → Server", "Gửi yêu cầu", "LOGIN, INVITE, INVITE_REPLY, CAR_STATE, FINISH, QUIT_MATCH, REMATCH_REPLY, LEADERBOARD_REQ."],
    ["Game Server", "Điều phối và trọng tài", "Quản lý kết nối, danh sách online, lời mời (timeout 30 s), phòng đua, countdown, đồng bộ vị trí, luật thắng thua, phát hiện mất kết nối."],
    ["Mũi tên Server → Client", "Phản hồi và đồng bộ", "LOGIN_RESULT, ONLINE_LIST, INVITE_RECEIVED, MATCH_START, COUNTDOWN, RACE_UPDATE, MATCH_RESULT, REMATCH_ASK, LEADERBOARD."],
    ["MySQL", "Lưu dữ liệu lâu dài", "Bảng players (tài khoản, điểm, thắng/thua/hòa), matches (kết quả trận), match_events (diễn biến trận)."],
], [1.55, 1.35, 3.37], caption="Vai trò các thành phần trong kiến trúc", first_col_bold=True)

table(["Thành phần", "Công nghệ", "Ghi chú"], [
    ["Ngôn ngữ", "Java 17 (JDK 17 LTS)", "Build bằng Maven, hai module client và server dùng chung module common (Message, model)."],
    ["Giao diện client", "Java Swing", "JFrame, JTable cho sảnh; JPanel tự vẽ bằng Java2D (paintComponent) cho màn hình đua kiểu game đua xe cổ điển nhìn từ trên xuống; javax.swing.Timer 50 ms cho vòng lặp và 16 ms cho vẽ lại."],
    ["Mạng", "java.net.Socket / ServerSocket", "TCP; ObjectInputStream / ObjectOutputStream; cổng mặc định 5000."],
    ["Đa luồng", "Thread, ExecutorService", "Mỗi client một luồng nhận và một luồng ghi (hàng đợi gửi); ScheduledExecutorService cho tick trận đấu và hẹn giờ lời mời."],
    ["Cơ sở dữ liệu", "MySQL 8 + JDBC", "Driver mysql-connector-j; PreparedStatement; transaction khi lưu kết quả và cập nhật điểm."],
], [1.35, 1.9, 3.02], caption="Công nghệ sử dụng", first_col_bold=True)

# ================================================================ 3. THIẾT KẾ CHUNG
heading("3. Thiết kế chung của hệ thống")
heading("3.1. Luồng hoạt động tổng thể", 2)
para("Sau khi đăng nhập thành công, người chơi xem danh sách người đang online kèm điểm và trạng thái Rảnh hoặc "
     "Đang thi đấu. Người chơi A chọn B để thách đấu. Server chuyển lời mời đến B và chờ tối đa 30 giây; B có thể "
     "chấp nhận hoặc từ chối, hết thời gian thì lời mời tự hủy và A được thông báo.")
para("Nếu B chấp nhận, server tạo phòng đua, chuyển cả hai client sang màn hình đua và đếm ngược 3, 2, 1. Hai "
     "người cùng thấy hai đường đua song song có cùng chiều dài và cùng một dòng xe cộ do server sinh ra theo seed. "
     "Trong lúc đua, client gửi trạng thái xe (vị trí, làn, tốc độ) lên server 20 lần mỗi giây; server kiểm tra "
     "hợp lệ, cập nhật vị trí hai xe rồi phát RACE_UPDATE cho cả hai. Khi một xe chạm vạch đích, server xác định "
     "kết quả, lưu vào CSDL, cập nhật điểm và hỏi hai bên có muốn thi đấu tiếp.")
figure("fig2_sequence.png", "Biểu đồ tuần tự từ đăng nhập đến kết thúc trận đấu")

heading("3.2. Giao thức trao đổi thông điệp", 2)
para("Mọi thông điệp giữa client và server đều là đối tượng Message cài đặt Serializable, gồm kiểu thông điệp "
     "(enum MessageType), dữ liệu kèm theo (payload) và thời điểm gửi. Cách này giúp hai phía chỉ cần một vòng "
     "lặp readObject / writeObject duy nhất và dễ mở rộng thêm loại thông điệp mới.")
code([
    "public class Message implements Serializable {",
    "    private static final long serialVersionUID = 1L;",
    "    private final MessageType type;     // LOGIN, INVITE, CAR_STATE, ...",
    "    private final Object payload;       // DTO tương ứng với từng loại",
    "    private final long timestamp = System.currentTimeMillis();",
    "",
    "    public Message(MessageType type, Object payload) {",
    "        this.type = type; this.payload = payload;",
    "    }",
    "    public MessageType getType() { return type; }",
    "    @SuppressWarnings(\"unchecked\")",
    "    public <T> T getPayload() { return (T) payload; }",
    "}",
])
table(["Loại thông điệp", "Hướng", "Dữ liệu kèm theo", "Ý nghĩa"], [
    ["LOGIN", "C → S", "username, password", "Yêu cầu đăng nhập (YC01)."],
    ["LOGIN_RESULT", "S → C", "ok, message, PlayerInfo", "Kết quả đăng nhập; kèm điểm hiện có nếu thành công."],
    ["ONLINE_LIST", "S → C", "List<PlayerInfo>", "Danh sách online: tên, điểm, thắng, trạng thái (YC02). Gửi lại mỗi khi danh sách đổi và gửi riêng cho người vừa đăng nhập ngay sau LOGIN_RESULT."],
    ["INVITE", "C → S", "targetUsername", "A thách đấu B (YC03)."],
    ["INVITE_RECEIVED", "S → C", "fromUsername, inviteId", "Server chuyển lời mời tới B; B hiện hộp thoại 30 giây."],
    ["INVITE_REPLY", "C → S", "inviteId, accept", "B chấp nhận hoặc từ chối (YC04)."],
    ["INVITE_RESULT", "S → C", "inviteId, status", "Báo cho A: REJECTED, TIMEOUT hoặc BUSY."],
    ["MATCH_START", "S → C", "roomId, opponent, trackSeed, obstacles", "Vào phòng đua; cả hai nhận cùng danh sách xe cộ, mỗi xe gồm làn, vị trí lúc xuất phát, tốc độ, kiểu xe (YC05, YC06)."],
    ["COUNTDOWN", "S → C", "value (3, 2, 1, 0)", "Đếm ngược; giá trị 0 nghĩa là bắt đầu đua."],
    ["CAR_STATE", "C → S", "distance, lane, speed", "Trạng thái xe của người chơi, gửi 20 lần/giây (YC07)."],
    ["RACE_UPDATE", "S → C", "RaceState hai xe, tick", "Vị trí, làn, tốc độ, quãng đường của cả hai xe do server đồng bộ; tick dùng để client tính vị trí xe cộ."],
    ["FINISH", "C → S", "clientTime", "Client báo xe đã qua vạch đích; server đối chiếu quãng đường trước khi công nhận."],
    ["MATCH_RESULT", "S → C", "winner, reason, điểm mới", "Kết quả trận: WIN, LOSE hoặc DRAW cùng lý do (FINISH, DRAW, QUIT, DISCONNECT) (YC08)."],
    ["REMATCH_ASK", "S → C", "roomId", "Hỏi có thi đấu tiếp không (YC09)."],
    ["REMATCH_REPLY", "C → S", "roomId, agree", "Trả lời thi đấu tiếp; cả hai đồng ý mới tạo ván mới."],
    ["QUIT_MATCH", "C → S", "roomId", "Chủ động thoát trận (YC10)."],
    ["PING / PONG", "C ⇄ S", "—", "Heartbeat mỗi 5 giây để phát hiện mất kết nối (YC11)."],
    ["LEADERBOARD_REQ / LEADERBOARD", "C ⇄ S", "List<RankRow>", "Xem bảng xếp hạng (YC12)."],
    ["LOGOUT", "C → S", "—", "Đăng xuất, server xóa khỏi danh sách online."],
], [1.75, 0.6, 1.6, 2.32], caption="Danh sách thông điệp của giao thức", first_col_bold=True, size=11, align_center_cols=(1,))

heading("3.3. Vòng đời phòng đua tại server", 2)
para("Mỗi phòng đua (Room) tại server là một máy trạng thái. Phòng được tạo khi lời mời được chấp nhận, đếm ngược "
     "ba giây, chuyển sang trạng thái đua, kết thúc khi có kết quả và chỉ quay lại ván mới khi cả hai người chơi "
     "đồng ý. Mọi chuyển trạng thái đều do server thực hiện nên hai client không thể lệch nhau.")
figure("fig7_states.png", "Máy trạng thái của một phòng đua")
table(["Trạng thái", "Điều kiện vào", "Server làm gì"], [
    ["WAITING", "Lời mời được chấp nhận", "Tạo Room, đặt cả hai người chơi sang trạng thái Đang thi đấu, sinh danh sách xe cộ theo seed, gửi MATCH_START."],
    ["COUNTDOWN", "Cả hai client báo đã vào phòng", "Gửi COUNTDOWN 3, 2, 1 cách nhau 1 giây; chưa nhận CAR_STATE."],
    ["RACING", "Đếm ngược về 0", "Chạy tick 50 ms: nhận CAR_STATE, kiểm tra hợp lệ, tính va chạm với xe cộ tại vị trí của tick đó, phát RACE_UPDATE, ghi match_events."],
    ["FINISHED", "Về đích, hòa, thoát hoặc mất kết nối", "Chốt kết quả, lưu matches, cập nhật điểm, gửi MATCH_RESULT rồi REMATCH_ASK."],
    ["CLOSED", "Một bên từ chối hoặc rời phòng", "Trả hai người chơi về trạng thái Rảnh, phát lại ONLINE_LIST, hủy Room."],
], [1.35, 1.6, 3.32], caption="Các trạng thái của phòng đua", first_col_bold=True)

heading("3.4. Luật thắng thua và tính điểm", 2)
para("Server là trọng tài duy nhất. Kết quả được xác định tại tick mà một xe có quãng đường lớn hơn hoặc bằng chiều "
     "dài đường đua. Nếu trong cùng một tick (50 ms) cả hai xe cùng đạt điều kiện thì tính hòa. Điểm được cộng "
     "trong một transaction cùng với việc lưu trận để tránh sai lệch khi có lỗi giữa chừng.")
para("Chướng ngại vật là 48 xe cộ chạy cùng chiều, mỗi làn 16 xe. Mỗi làn nhận một tốc độ cố định lấy từ bộ "
     "110, 160 và 210 km/h, xáo thứ tự theo seed của ván, nên các xe cùng làn không bao giờ đè lên nhau; xe của "
     "người chơi chạy tối đa 360 km/h nên luôn vượt được. Vị trí một xe cộ ở tick k tính theo công thức "
     "vị trí lúc xuất phát + tốc độ × k × 50 ms. Server (khi xét va chạm) và client (khi vẽ) cùng gọi "
     "Obstacle.positionAt(tick) nên hai bên luôn khớp mà không phải gửi vị trí từng xe cộ qua mạng. Va chạm "
     "được xét tại vị trí xe cộ đã di chuyển: cùng làn và chồng lên nhau thì tốc độ xe về 0 trong 1,5 giây. Mỗi "
     "xe cộ chỉ bị đâm một lần; trên màn hình xe bị đâm nổ và biến mất.")
table(["Tình huống", "Kết quả", "Điểm người A", "Điểm người B", "end_reason"], [
    ["A về đích trước B", "A thắng", "+1", "0", "FINISH"],
    ["A và B về đích trong cùng một tick", "Hòa", "+1", "+1", "DRAW"],
    ["A chủ động thoát trận", "B thắng", "0", "+1", "QUIT"],
    ["A mất kết nối quá timeout (15 giây)", "B thắng", "0", "+1", "DISCONNECT"],
    ["Cả hai cùng mất kết nối", "Hủy trận", "0", "0", "ABORTED"],
], [2.1, 1.0, 0.95, 0.95, 1.27], caption="Luật tính điểm", align_center_cols=(1, 2, 3, 4))

heading("3.5. Xử lý các tình huống đặc biệt", 2)
table(["Tình huống", "Xử lý của server"], [
    ["B từ chối lời mời", "Gửi INVITE_RESULT(REJECTED) cho A; không tạo phòng; cả hai vẫn Rảnh."],
    ["B không phản hồi trong 30 giây", "Hẹn giờ trong InviteManager hết hạn, hủy lời mời, gửi INVITE_RESULT(TIMEOUT) cho A và đóng hộp thoại ở B."],
    ["A mời B nhưng B vừa vào trận khác", "Trả INVITE_RESULT(BUSY); danh sách online được gửi lại để A thấy trạng thái mới."],
    ["Một người thoát trận", "Gửi MATCH_RESULT cho đối thủ với reason = QUIT, xử người thoát thua, lưu trận, đóng phòng."],
    ["Mất kết nối quá timeout", "Không nhận PING trong 15 giây hoặc readObject ném IOException: xử thua, lưu end_reason = DISCONNECT, báo đối thủ."],
    ["Client gửi CAR_STATE bất thường", "Server giới hạn tốc độ tối đa và quãng đường tăng tối đa mỗi tick; giá trị vượt ngưỡng bị cắt về ngưỡng."],
    ["Cả hai chọn thi đấu tiếp", "Tạo ván mới trong cùng phòng: sinh lại dòng xe cộ với seed mới, đặt lại quãng đường bằng 0, đếm ngược lại."],
    ["Một client ngừng nhận dữ liệu (treo, mạng nghẽn)", "Hàng đợi gửi của client đó đầy (256 thông điệp, khoảng 12 giây RACE_UPDATE) thì server ngắt kết nối và xử lý như mất kết nối; các phòng khác không bị ảnh hưởng."],
    ["Một trong hai từ chối thi đấu tiếp", "Đóng phòng, cả hai trở về sảnh với trạng thái Rảnh."],
], [2.0, 4.27], caption="Xử lý tình huống đặc biệt", first_col_bold=True)

# ================================================================ 4. CÁ NHÂN 1
heading("4. Phần cá nhân 1 – Server và kiến trúc (Phạm Thị Thu Phương)")
para("Phạm Thị Thu Phương (thành viên 1) thực hiện phần server (task S1–S7). Server nhận nhiều kết nối TCP, tạo ClientHandler cho từng client, quản "
     "lý danh sách online, lời mời, phòng đua và xác định kết quả trận đấu. Server đọc ghi dữ liệu với MySQL bằng "
     "JDBC thông qua lớp DAO. Toàn bộ trạng thái dùng chung (danh sách online, lời mời, phòng) được giữ trong các "
     "cấu trúc thread-safe để nhiều luồng ClientHandler truy cập đồng thời không gây lỗi.")
figure("fig3_server.png", "Cấu trúc bên trong Game Server", width_in=6.0)
table(["Thành phần trên hình", "Chức năng", "Mô tả"], [
    ["ServerSocket", "Nhận kết nối", "Mở cổng 5000, vòng lặp accept(); mỗi Socket mới được bọc trong một ClientHandler và giao cho ExecutorService."],
    ["ClientHandler", "Xử lý từng client", "Luồng đọc Message bằng ObjectInputStream, chuyển cho bộ xử lý theo type. Gửi trả qua hàng đợi riêng (tối đa 256 thông điệp) và một luồng ghi ObjectOutputStream riêng, nên send() không bao giờ chặn."],
    ["SessionManager", "Người chơi online", "ConcurrentHashMap<username, ClientHandler>; trạng thái FREE / IN_MATCH; phát ONLINE_LIST khi có thay đổi; kiểm tra heartbeat."],
    ["InviteManager", "Lời mời thách đấu", "Lưu lời mời đang chờ, hẹn giờ 30 giây bằng ScheduledExecutorService, xử lý accept / reject / timeout."],
    ["RoomManager / MatchService", "Trọng tài trận đấu", "Tạo Room, sinh xe cộ (ObstacleGenerator), countdown, tick 50 ms, xét va chạm (CarSim), đồng bộ RACE_UPDATE, xét về đích, hòa, thoát, mất kết nối, hỏi thi đấu tiếp."],
    ["PlayerDAO / MatchDAO", "Truy cập CSDL", "Đăng nhập, lấy bảng xếp hạng, lưu trận và cập nhật điểm trong transaction."],
], [1.55, 1.25, 3.47], caption="Thành phần của server", first_col_bold=True)

heading("4.1. Các lớp chính phía server", 2)
table(["Lớp", "Gói", "Trách nhiệm chính"], [
    ["GameServer", "server", "main(); mở ServerSocket; vòng lặp accept; khởi tạo các manager."],
    ["ClientHandler", "server.net", "implements Runnable; vòng lặp đọc Message; gửi Message; đóng kết nối và báo SessionManager khi lỗi."],
    ["SessionManager", "server.core", "Đăng nhập / đăng xuất, danh sách online, trạng thái người chơi, heartbeat."],
    ["InviteManager", "server.core", "Tạo, hủy, hết hạn lời mời; kiểm tra người được mời đang Rảnh."],
    ["RoomManager", "server.core", "Tạo và hủy Room; ánh xạ người chơi → Room."],
    ["Room", "server.core", "Trạng thái một trận: hai CarSim, danh sách xe cộ, máy trạng thái, tick."],
    ["CarSim", "server.core", "Trọng tài từng xe: cắt tốc độ, tính quãng đường theo tick, xét va chạm với Obstacle.positionAt(tick)."],
    ["ObstacleGenerator", "server.core", "Sinh 48 xe cộ theo seed: mỗi làn một tốc độ, hai xe cùng làn cách nhau ít nhất 40 m."],
    ["MatchService", "server.core", "Luật thắng thua, tính điểm, gọi DAO lưu kết quả."],
    ["AccountService", "server.core", "Đăng ký tài khoản, bảng xếp hạng, lịch sử trận."],
    ["PlayerDAO, MatchDAO", "server.db", "JDBC với PreparedStatement; DbConnection quản lý kết nối."],
    ["Message, MessageType, các DTO", "common", "Dùng chung cho client và server (PlayerInfo, RaceState, Obstacle, MatchResult…); do Trần Việt Anh chốt ở phần nền (D2)."],
], [1.7, 1.1, 3.47], caption="Các lớp phía server", first_col_bold=True, size=11)

heading("4.2. Vòng lặp xử lý của ClientHandler", 2)
code([
    "public void run() {",
    "    try (ObjectInputStream in = new ObjectInputStream(socket.getInputStream())) {",
    "        socket.setSoTimeout(15_000);      // không nhận gì trong 15 s → mất kết nối",
    "        while (running) {",
    "            Message msg = (Message) in.readObject();",
    "            switch (msg.getType()) {",
    "                case LOGIN        -> sessionManager.login(this, msg.getPayload());",
    "                case INVITE       -> inviteManager.invite(player, msg.getPayload());",
    "                case INVITE_REPLY -> inviteManager.reply(player, msg.getPayload());",
    "                case CAR_STATE    -> roomManager.onCarState(player, msg.getPayload());",
    "                case FINISH       -> roomManager.onFinish(player);",
    "                case QUIT_MATCH   -> roomManager.onQuit(player);",
    "                case REMATCH_REPLY-> roomManager.onRematch(player, msg.getPayload());",
    "                case PING         -> send(new Message(MessageType.PONG, null));",
    "                default           -> log.warn(\"Unknown message {}\", msg.getType());",
    "            }",
    "        }",
    "    } catch (IOException | ClassNotFoundException e) {",
    "        // SocketTimeoutException cũng rơi vào đây",
    "        roomManager.onDisconnect(player);   // xử thua nếu đang trong trận",
    "        sessionManager.logout(player);      // gỡ khỏi danh sách online",
    "    }",
    "}",
])
para("Phát hiện mất kết nối (YC11) dựa trên hai cơ chế bổ sung nhau: client gửi PING mỗi 5 giây, và socket phía "
     "server đặt SoTimeout 15 giây. Khi client tắt ứng dụng, readObject ném IOException ngay; khi client mất mạng "
     "im lặng, SocketTimeoutException xuất hiện sau 15 giây. Cả hai trường hợp cùng đi vào nhánh xử lý "
     "onDisconnect, nên logic xử thua chỉ viết một lần.")
para("Về đồng bộ đa luồng: mỗi Room có khóa riêng (synchronized trên đối tượng Room) để hai ClientHandler của "
     "hai người chơi và luồng tick không đồng thời sửa trạng thái xe. Việc gửi Message được tách khỏi luồng gọi: "
     "send() chỉ đưa thông điệp vào hàng đợi của client, một luồng ghi riêng lấy ra và ghi vào ObjectOutputStream. "
     "Luồng tick dùng chung giữa các phòng, nên nếu ghi trực tiếp thì một client ngừng nhận (bộ đệm TCP đầy) sẽ "
     "làm cả phòng đứng hình; với hàng đợi, chỉ client đó bị ngắt khi hàng đợi đầy.")
heading("4.3. Gửi thông điệp không chặn", 2)
code([
    "private final BlockingQueue<Message> outbox = new LinkedBlockingQueue<>(256);",
    "",
    "public void send(Message m) {             // gọi từ luồng tick, hẹn giờ, client khác",
    "    if (!outbox.offer(m)) closeSocket();   // không nhận kịp → ngắt, xử như mất kết nối",
    "}",
    "",
    "private void writeLoop() {                // luồng ghi riêng của mỗi client",
    "    while (!socket.isClosed()) {",
    "        out.writeObject(outbox.take());",
    "        out.reset();                       // không gửi lại bản cũ của DTO",
    "        out.flush();",
    "    }",
    "}",
])
para("Khi một người đăng nhập, SessionManager phát ONLINE_LIST cho mọi người trước khi ClientHandler trả "
     "LOGIN_RESULT; client chỉ mở sảnh sau LOGIN_RESULT nên bản danh sách đó đến quá sớm. Vì vậy ClientHandler "
     "gửi thêm ONLINE_LIST cho chính người vừa vào ngay sau LOGIN_RESULT để sảnh có danh sách ngay.")
para("Server nhận tham số cổng và tùy chọn --memory. Với --memory, server dùng kho dữ liệu trong bộ nhớ (6 tài "
     "khoản mẫu, mất khi tắt server) thay cho MySQL, phục vụ chạy thử nhanh và demo; mặc định server dùng MySQL.")

# ================================================================ 5. CÁ NHÂN 2
heading("5. Phần cá nhân 2 – Client: đăng nhập, sảnh chờ và thách đấu (Vũ Văn Hiếu)")
para("Vũ Văn Hiếu (thành viên 2) thực hiện phần client trước khi vào trận (task C1–C7): màn hình đăng nhập, sảnh chờ hiển thị danh sách "
     "người chơi online, gửi lời mời thách đấu và hộp thoại nhận lời mời có đếm ngược 30 giây. Client được tổ "
     "chức theo MVC: View là các JFrame / JDialog, Controller là lớp NetworkClient chạy một luồng nhận riêng, "
     "Model là dữ liệu PlayerInfo nhận từ server. Mọi cập nhật giao diện đều được đẩy về luồng Swing bằng "
     "SwingUtilities.invokeLater để không vi phạm quy tắc một luồng của Swing.")
figure("fig4_lobby.png", "Sảnh chờ và hộp thoại nhận lời mời thách đấu")
table(["Thành phần trên hình", "Chức năng", "Mô tả"], [
    ["Thanh tiêu đề", "Thông tin phiên", "Tên người chơi đang đăng nhập và tổng điểm hiện có, cập nhật sau mỗi trận."],
    ["Bảng danh sách online", "Hiển thị YC02", "JTable với các cột Tên, Điểm, Thắng, Trạng thái; dữ liệu là ONLINE_LIST mới nhất từ server."],
    ["Nút Thách đấu", "Gửi INVITE (YC03)", "Chỉ hiện với người Rảnh; click gửi INVITE rồi khóa nút và hiện trạng thái Đang chờ B trả lời."],
    ["Hộp thoại Lời mời thách đấu", "Nhận INVITE_RECEIVED (YC04)", "Hiện ở máy người được mời; thanh tiến trình và số giây còn lại giảm dần từ 30; nút Chấp nhận / Từ chối."],
    ["Nút Bảng xếp hạng", "Gửi LEADERBOARD_REQ (YC12)", "Mở cửa sổ xếp hạng (LeaderboardFrame, phần cá nhân 4)."],
    ["Nút Lịch sử trận / Đăng xuất", "Tiện ích", "Mở lịch sử trận của mình (HistoryFrame, phần cá nhân 4); gửi LOGOUT và quay về màn hình đăng nhập."],
], [1.5, 1.75, 3.02], caption="Thành phần giao diện sảnh chờ", first_col_bold=True)

heading("5.1. Các lớp chính phía client", 2)
table(["Lớp", "Vai trò MVC", "Trách nhiệm chính"], [
    ["LoginFrame", "View", "Nhập username, password; gửi LOGIN; hiện lỗi nếu LOGIN_RESULT thất bại."],
    ["LobbyFrame", "View", "JTable danh sách online; nút Thách đấu, Bảng xếp hạng, Lịch sử trận, Đăng xuất."],
    ["InviteDialog", "View", "JDialog modal hiển thị lời mời với javax.swing.Timer đếm ngược 30 giây; tự đóng khi hết giờ hoặc nhận INVITE_RESULT."],
    ["LeaderboardFrame, HistoryFrame, RegisterDialog", "View", "Bảng xếp hạng, lịch sử trận, đăng ký tài khoản; do Mai Anh làm (M2–M4), Hiếu nối nút mở từ sảnh và màn hình đăng nhập (C5)."],
    ["NetworkClient", "Controller", "Mở Socket, luồng nhận readObject, gửi Message; phân phối thông điệp tới màn hình đang mở; PING mỗi 5 giây."],
    ["ClientState", "Model", "Người chơi hiện tại, danh sách PlayerInfo, lời mời đang chờ, roomId hiện tại."],
], [1.5, 1.05, 3.72], caption="Các lớp phía client (phần sảnh)", first_col_bold=True, size=11)

heading("5.2. Xử lý lời mời thách đấu ở client", 2)
code([
    "// Nhận INVITE_RECEIVED từ NetworkClient (đã ở luồng Swing)",
    "void onInviteReceived(InviteInfo info) {",
    "    InviteDialog dlg = new InviteDialog(lobbyFrame, info.from(), 30);",
    "    dlg.onAccept(() -> net.send(new Message(MessageType.INVITE_REPLY,",
    "                                            new InviteReply(info.id(), true))));",
    "    dlg.onReject(() -> net.send(new Message(MessageType.INVITE_REPLY,",
    "                                            new InviteReply(info.id(), false))));",
    "    dlg.onTimeout(dlg::dispose);   // server cũng tự hủy sau 30 s, client chỉ đóng hộp thoại",
    "    dlg.setVisible(true);",
    "}",
])
para("Điểm quan trọng là thời hạn 30 giây được tính tại server (InviteManager) chứ không phụ thuộc đồng hồ của "
     "client; bộ đếm trên hộp thoại chỉ để người chơi thấy thời gian còn lại. Nếu người được mời bấm Chấp nhận "
     "sau khi server đã hủy, server trả INVITE_RESULT(TIMEOUT) và client hiển thị thông báo lời mời đã hết hạn.")

# ================================================================ 6. CÁ NHÂN 3
heading("6. Phần cá nhân 3 – Nền tảng và giao diện màn hình đua (Trần Việt Anh)")
para("Trần Việt Anh (thành viên 3) làm phần nền tảng (task D1–D7: Maven nhiều module, module common với Message, "
     "MessageType, DTO và GameConfig, lược đồ CSDL cùng DAO nền, CI, công cụ racing-tools, kiểm thử tích hợp, đóng "
     "gói) và giao diện màn hình đua (task R1–R6). Màn hình đua theo phong cách game đua xe cổ điển nhìn từ trên "
     "xuống: màn hình chia đôi, bên trái là đường của mình, bên phải là đường của đối thủ, mỗi bên có một camera "
     "cuộn bám theo xe. Toàn bộ hình ảnh (đường, cảnh quan, xe, vụ nổ, chữ) được vẽ bằng Java2D dạng pixel, "
     "không dùng file ảnh. Người chơi chỉ điều khiển xe của mình; dữ liệu xe đối thủ do server đồng bộ qua "
     "RACE_UPDATE.")
figure("fig5_race.png", "Màn hình đua: bên trái vừa đâm xe cộ, bên phải đối thủ sắp về đích")
table(["Thành phần trên hình", "Chức năng", "Mô tả"], [
    ["(1) Thanh HUD đen", "Chỉ số mỗi bên", "1P / 2P kèm tên, tốc độ KM/H, quãng đường dạng 1848/2500 M và trạng thái 1ST, 2ND, CRASH! hoặc FINISH!."],
    ["(2) Minimap ở mép ngoài", "Tiến độ", "Dải dọc biểu diễn 2500 m, chấm đỏ là xe mình, chấm xanh là đối thủ, cho biết ai đang dẫn."],
    ["(3) Xe của mình, vụ nổ", "Điều khiển", "Xe pixel đỏ trắng đứng cố định gần đáy khung, đường cuộn theo tốc độ. Đổi làn trượt mượt theo A / D, tăng giảm tốc theo W / S. Khi va chạm hiện vụ nổ, xe nhấp nháy, tốc độ về 0 trong 1,5 giây."],
    ["(4) Xe cộ", "Chướng ngại vật", "48 xe, 6 kiểu (sedan nhiều màu, xe tải, xe van) theo kind; vị trí Obstacle.positionAt(tick) nội suy giữa hai RACE_UPDATE. Xe bị đâm nổ và biến mất."],
    ["(5) Cảnh quan", "Cảm giác tốc độ", "Cỏ, cây thông, bụi, nhà mái đỏ, ao hai bên đường; sinh cố định theo hàm băm vị trí nên khi cuộn không nhấp nháy."],
    ["(6) Vạch START / FINISH", "Mốc đường đua", "Hiện khi lọt vào khung nhìn; xe về đích khi quãng đường ≥ 2500 m."],
    ["(7) Đường bên phải", "Xe đối thủ", "Camera riêng bám xe đối thủ, chỉ hiển thị theo RACE_UPDATE, không nhận phím."],
    ["(8) Dòng trạng thái", "Hướng dẫn", "Phím điều khiển và trạng thái trận. Trước khi đua, hộp đếm ngược kiểu START (3, 2, 1, GO) phủ giữa màn hình; nút Thoát trận (Esc) gửi QUIT_MATCH (YC10)."],
], [1.55, 1.3, 3.42], caption="Thành phần giao diện màn hình đua", first_col_bold=True)

heading("6.1. Các lớp của màn hình đua", 2)
table(["Lớp", "Trách nhiệm chính"], [
    ["RaceFrame", "Cửa sổ trận: nhận COUNTDOWN, RACE_UPDATE, MATCH_RESULT, REMATCH_ASK; vòng lặp 50 ms gửi CAR_STATE; xử lý phím và Thoát trận."],
    ["RacePanel", "Chia đôi màn hình, HUD, minimap, hộp đếm ngược, dòng trạng thái; Timer 16 ms vẽ lại."],
    ["TrackView", "Camera cuộn của một bên: nền cỏ và mặt đường, cảnh quan, vạch làn, START / FINISH, xe cộ, xe người chơi, vụ nổ."],
    ["PixelArt", "Dựng sẵn hình pixel (xe theo kind, cây, nhà, ao, vụ nổ, chữ pixel) vào bộ đệm ảnh để mỗi khung hình chỉ việc vẽ lại."],
    ["CarModel", "Trạng thái một xe phía client: dự đoán cục bộ; xe mình nhận quãng đường, va chạm, về đích từ server (applyServerOwn), xe đối thủ ghi đè toàn bộ (applyServer)."],
    ["ResultDialog, RematchDialog", "Hộp kết quả (thắng, thua, hòa, lý do, điểm mới) và hỏi thi đấu tiếp."],
    ["RaceDemo", "Chạy màn hình đua không cần server: giả lập đếm ngược, đối thủ tự lái, xe cộ chạy, va chạm, kết quả."],
], [1.75, 4.52], caption="Các lớp phía client (màn hình đua)", first_col_bold=True, size=11)

heading("6.2. Điều khiển và vòng lặp trò chơi", 2)
table(["Phím", "Hành động", "Mô tả"], [
    ["Giữ W hoặc ↑", "Tăng tốc", "Tốc độ tăng đều 120 km/h mỗi giây khi giữ phím, tối đa 360 km/h (từ 0 lên tối đa trong 3 giây)."],
    ["Giữ S hoặc ↓", "Phanh", "Tốc độ giảm 300 km/h mỗi giây khi giữ phím, tối thiểu 0."],
    ["Không giữ W", "Nhả ga", "Xe tự giảm chậm 40 km/h mỗi giây."],
    ["A hoặc ←", "Sang làn trái", "Chuyển làn 1 → 0 hoặc 2 → 1; không đổi nếu đang ở làn ngoài cùng."],
    ["D hoặc →", "Sang làn phải", "Chuyển làn 0 → 1 hoặc 1 → 2."],
    ["Esc", "Thoát trận", "Tương đương nút Thoát trận."],
], [1.1, 1.45, 3.72], caption="Phím điều khiển", first_col_bold=True, align_center_cols=(0,))
para("Phím tăng tốc và phanh là phím giữ: client ghi nhận lúc nhấn và lúc nhả (bỏ qua lặp phím của hệ điều "
     "hành), tốc độ đổi dần theo thời gian giữ chứ không theo số lần nhấn. Vòng lặp trò chơi ở client có hai "
     "nhịp. Timer 50 ms đổi tốc độ theo phím đang giữ, cập nhật quãng đường và gửi "
     "CAR_STATE lên server. Timer 16 ms (khoảng 60 khung hình mỗi giây) chỉ vẽ lại để đường cuộn mượt. Khi nhận "
     "RACE_UPDATE, client lấy quãng đường, va chạm và về đích của xe mình theo server nhưng giữ làn và tốc độ theo "
     "phím đang bấm (gói RACE_UPDATE được tính trước khi server nhận CAR_STATE mới nhất, nếu ghi đè làn thì lần "
     "đổi làn vừa bấm sẽ bị bật ngược); xe đối thủ ghi đè toàn bộ theo server. Client lưu lại số tick; giữa hai lần cập nhật, "
     "TrackView nội suy thời gian để tính vị trí xe cộ bằng cùng công thức Obstacle.positionAt với server. Khi "
     "server báo xe vừa bị choáng, client tìm xe cộ cùng làn đang chồng lên xe mình, cho nó nổ và bỏ khỏi đường.")
code([
    "tick = new Timer(GameConfig.TICK_MS, e -> onTick());   // 50 ms: tiến xe, gửi CAR_STATE",
    "",
    "private void onRaceUpdate(RaceState s) {",
    "    panel.setRaceTick(s.tick());                  // mốc để vẽ xe cộ",
    "    if (myCar.applyServerOwn(s.me())) {           // giữ làn, tốc độ; server báo va chạm",
    "        panel.flashMine();                        // nổ + nhấp nháy",
    "    }",
    "    if (opponentCar.applyServer(s.opponent())) {",
    "        panel.flashOpponent();",
    "    }",
    "}",
])

# ================================================================ 7. CÁ NHÂN 4
heading("7. Phần cá nhân 4 – Dữ liệu: bảng xếp hạng, đăng ký, lịch sử trận (Nguyễn Trần Mai Anh)")
para("Nguyễn Trần Mai Anh (thành viên 4) phụ trách phần dữ liệu (task M1–M6): mô tả CSDL trong docs/DB.md, bảng "
     "xếp hạng đầu cuối, đăng ký tài khoản, lịch sử trận, ghi diễn biến vào match_events và kiểm tra dữ liệu bằng SQL "
     "khi kiểm thử tích hợp. Lược đồ và DAO nền do Trần Việt Anh dựng ở phần nền tảng (D3, D4). Khi trận kết thúc, server lưu kết quả vào bảng "
     "matches, ghi các sự kiện quan trọng vào match_events và cập nhật điểm, số trận thắng, thua, hòa trong bảng "
     "players. Bảng xếp hạng được lấy trực tiếp từ players, sắp xếp theo tổng điểm giảm dần rồi tổng số trận thắng "
     "giảm dần (YC12).")
figure("fig6_erd.png", "Mô hình dữ liệu quan hệ")
table(["Thành phần trên hình", "Chức năng", "Mô tả"], [
    ["players", "Lưu người chơi", "Tài khoản (username duy nhất, mật khẩu băm), tổng điểm, số trận thắng, thua, hòa. Điểm và số thắng là cột lưu sẵn để truy vấn xếp hạng nhanh."],
    ["matches", "Lưu kết quả trận", "Mã phòng, hai người chơi, người thắng, trạng thái, lý do kết thúc, thời gian bắt đầu và kết thúc."],
    ["Ba quan hệ players → matches", "Hai người chơi và người thắng", "player1_id, player2_id, winner_id cùng là khóa ngoại tới players; winner_id NULL khi hòa hoặc hủy trận."],
    ["match_events", "Lưu diễn biến", "Sự kiện theo thời gian: START, COLLISION, FINISH, QUIT, DISCONNECT, REMATCH với payload JSON; COLLISION ghi làn, quãng đường của xe và vị trí xe cộ lúc va chạm."],
    ["Truy vấn xếp hạng", "YC12", "ORDER BY points DESC, wins DESC; có chỉ mục (points, wins) để sắp xếp nhanh."],
], [1.6, 1.3, 3.37], caption="Thành phần của mô hình dữ liệu", first_col_bold=True)

heading("7.1. Lược đồ cơ sở dữ liệu", 2)
code([
    "CREATE TABLE players (",
    "    player_id     INT AUTO_INCREMENT PRIMARY KEY,",
    "    username      VARCHAR(50)  NOT NULL UNIQUE,",
    "    password_hash VARCHAR(255) NOT NULL,",
    "    points        INT NOT NULL DEFAULT 0,",
    "    wins          INT NOT NULL DEFAULT 0,",
    "    losses        INT NOT NULL DEFAULT 0,",
    "    draws         INT NOT NULL DEFAULT 0,",
    "    created_at    DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,",
    "    INDEX idx_rank (points DESC, wins DESC)",
    ");",
    "",
    "CREATE TABLE matches (",
    "    match_id    INT AUTO_INCREMENT PRIMARY KEY,",
    "    room_code   VARCHAR(20) NOT NULL,",
    "    player1_id  INT NOT NULL,",
    "    player2_id  INT NOT NULL,",
    "    winner_id   INT NULL,",
    "    status      ENUM('PLAYING','FINISHED','ABORTED') NOT NULL DEFAULT 'PLAYING',",
    "    end_reason  ENUM('FINISH','DRAW','QUIT','DISCONNECT','ABORTED') NULL,",
    "    started_at  DATETIME NOT NULL,",
    "    ended_at    DATETIME NULL,",
    "    FOREIGN KEY (player1_id) REFERENCES players(player_id),",
    "    FOREIGN KEY (player2_id) REFERENCES players(player_id),",
    "    FOREIGN KEY (winner_id)  REFERENCES players(player_id)",
    ");",
    "",
    "CREATE TABLE match_events (",
    "    event_id   BIGINT AUTO_INCREMENT PRIMARY KEY,",
    "    match_id   INT NOT NULL,",
    "    player_id  INT NOT NULL,",
    "    event_type VARCHAR(30) NOT NULL,",
    "    event_time DATETIME(3) NOT NULL,",
    "    payload    JSON NULL,",
    "    FOREIGN KEY (match_id)  REFERENCES matches(match_id),",
    "    FOREIGN KEY (player_id) REFERENCES players(player_id)",
    ");",
], size=9)

heading("7.2. Lưu kết quả và cập nhật điểm", 2)
para("Việc lưu kết quả và cộng điểm phải xảy ra cùng nhau, vì vậy MatchDAO thực hiện trong một transaction JDBC: "
     "tắt auto-commit, cập nhật bản ghi matches, cộng điểm cho từng người chơi rồi commit; nếu có lỗi thì rollback "
     "để không xảy ra tình trạng trận đã ghi mà điểm chưa cộng.")
code([
    "public void saveResult(MatchResult r) throws SQLException {",
    "    try (Connection c = DbConnection.get()) {",
    "        c.setAutoCommit(false);",
    "        String sqlMatch  = \"UPDATE matches SET winner_id=?, status='FINISHED', \"",
    "                         + \"end_reason=?, ended_at=NOW() WHERE match_id=?\";",
    "        String sqlPlayer = \"UPDATE players SET points=points+?, wins=wins+?, \"",
    "                         + \"losses=losses+?, draws=draws+? WHERE player_id=?\";",
    "        try (PreparedStatement m = c.prepareStatement(sqlMatch);",
    "             PreparedStatement p = c.prepareStatement(sqlPlayer)) {",
    "            m.setObject(1, r.winnerId());",
    "            m.setString(2, r.reason().name());",
    "            m.setInt(3, r.matchId());",
    "            m.executeUpdate();",
    "            for (PlayerDelta d : r.deltas()) {   // mỗi người chơi một dòng",
    "                p.setInt(1, d.points()); p.setInt(2, d.win());",
    "                p.setInt(3, d.loss());   p.setInt(4, d.draw());",
    "                p.setInt(5, d.playerId());",
    "                p.addBatch();",
    "            }",
    "            p.executeBatch();",
    "            c.commit();",
    "        } catch (SQLException e) { c.rollback(); throw e; }",
    "    }",
    "}",
], size=9)

heading("7.3. Truy vấn bảng xếp hạng", 2)
code([
    "SELECT RANK() OVER (ORDER BY points DESC, wins DESC) AS hang,",
    "       username, points, wins, losses, draws",
    "FROM players",
    "ORDER BY points DESC, wins DESC",
    "LIMIT 100;",
])
para("Client nhận danh sách RankRow qua thông điệp LEADERBOARD và hiển thị trong JTable với các cột Hạng, Tên, "
     "Điểm, Thắng, Thua, Hòa. Vì điểm và số trận thắng được cập nhật ngay trong transaction lưu kết quả, bảng "
     "xếp hạng luôn phản ánh đúng trận vừa kết thúc mà không cần tính lại từ bảng matches.")

# ================================================================ 8. PHÂN CÔNG
heading("8. Phân công công việc")
table(["Thành viên", "Phần phụ trách", "Sản phẩm bàn giao"], [
    ["Phạm Thị Thu Phương", "Server, S1–S7 (mục 4)", "GameServer, ClientHandler (hàng đợi gửi, luồng ghi), SessionManager, InviteManager, Room, RoomManager, CarSim, ObstacleGenerator, MatchService, AccountService; chịu tải 2 phòng song song."],
    ["Vũ Văn Hiếu", "Client sảnh, C1–C7 (mục 5)", "NetworkClient, ClientState, LoginFrame, LobbyFrame, InviteDialog; điều hướng sảnh ↔ trận, xử lý mất kết nối phía client."],
    ["Trần Việt Anh", "Nền tảng D1–D7, màn hình đua R1–R6 (mục 6)", "Maven, module common, schema.sql, seed.sql, DbConnection, PlayerDAO, MatchDAO nền, CI, racing-tools; RaceFrame, RacePanel, TrackView, PixelArt, CarModel, ResultDialog, RematchDialog, RaceDemo; dẫn kiểm thử tích hợp, đóng gói bản chạy."],
    ["Nguyễn Trần Mai Anh", "Dữ liệu SQL, M1–M6 (mục 7)", "docs/DB.md, LeaderboardFrame, RegisterDialog, HistoryFrame, MatchDAO.findRecentByPlayer, ghi match_events, kiểm tra SQL trong docs/TEST-PLAN.md."],
    ["Cả nhóm", "Tích hợp và kiểm thử", "Chạy kịch bản T1–T19 trong docs/TEST-PLAN.md trên 2 máy cùng mạng LAN (các tình huống ở mục 3.5), viết báo cáo."],
], [1.6, 1.9, 2.77], caption="Phân công công việc trong nhóm", first_col_bold=True)

# ================================================================ 9. KẾT LUẬN
heading("9. Kết luận")
para("Thiết kế bảo đảm Server là trọng tài trung tâm: mọi luật thắng thua, thời hạn lời mời, phát hiện mất kết "
     "nối và cập nhật điểm đều được quyết định tại server nên hai client không thể lệch kết quả. Server phục vụ "
     "nhiều client đồng thời nhờ mô hình một luồng cho mỗi kết nối cùng các cấu trúc dữ liệu thread-safe. Giao "
     "thức thông điệp thống nhất trên TCP giúp việc mở rộng thêm tính năng chỉ cần bổ sung loại thông điệp mới. "
     "Dữ liệu trận được lưu đầy đủ trong MySQL để phục vụ bảng xếp hạng và xem lại diễn biến, và từng phần cá "
     "nhân đều được ánh xạ rõ ràng với thành phần cụ thể trên các hình thiết kế.")

doc.save(OUT)
print("saved", OUT)
